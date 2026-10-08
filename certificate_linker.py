"""
certificate_linker.py  --  Coursera certificate-link filler (works for any semester)

INPUT
  1. Abstract workbook : one sheet per class (e.g. "I MCA", "III MCA") with columns
                         Roll No / External Id, Name, Course ... (header may be in any of first 10 rows)
  2. Coursera export(s): any number of "LearnerActivity ... CourseraEnterpriseExport" files.
                         Sheets with "Course Certificate URL" or "Specialization Certificate URL"
                         are used automatically; other sheets are ignored.

OUTPUT (new workbook)
  * One sheet per abstract sheet, same layout/formatting + a "Certificate link" column
  * Verification sheet : how every row was matched + name / roll-no cross-check
  * SUMMARY sheet      : course-wise counts (live formulas)

Matching rule : Course name (cleaned)  +  Roll No  and/or  Email  (Name only as last resort).
Why email too?: Coursera leaves "External ID" blank on many rows, so Roll No alone misses students.

Command line:
  python certificate_linker.py abstract.xlsx export1.xlsx export2.xlsx -o output.xlsx
"""
import argparse
import difflib
import re
from collections import defaultdict
from copy import copy

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NA_TEXT = "Link not available"

ID_ALIASES = {"external id", "roll no", "roll no.", "register no", "register no.",
              "reg no", "reg no.", "register number", "roll number", "usn"}
NAME_ALIASES = {"name", "student name"}
COURSE_ALIASES = {"course", "course name"}
EMAIL_ALIASES = {"email", "email id", "e-mail"}
LINK_ALIASES = {"certificate link", "certificate url", "certificate"}

RAW_SHEET_TYPES = [  # (name column, url column, completion-time column)
    ("course", "course certificate url", "completion time"),
    ("specialization", "specialization certificate url", "specialization completion time"),
]

# ----------------------------------------------------------------------------- helpers
def _blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ""


def norm(v):
    """lower-case, collapse spaces/tabs/nbsp, unify dashes."""
    if _blank(v):
        return ""
    s = str(v).replace("\xa0", " ").replace("\u2013", "-").replace("\u2014", "-")
    return re.sub(r"\s+", " ", s).strip().lower()


def norm_id(v):
    return re.sub(r"\s+", "", norm(v)).upper()


def name_key(v):
    return re.sub(r"[^a-z0-9]", "", norm(v))


def name_tokens(v):
    return frozenset(re.findall(r"[a-z0-9]+", norm(v)))


# ----------------------------------------------------------------------------- raw exports
def load_raw(files):
    """Return list of dict records from all Coursera export files."""
    recs = []
    for f in files:
        if hasattr(f, "seek"):
            f.seek(0)
        sheets = pd.read_excel(f, sheet_name=None, dtype=object)
        for sheet_name, df in sheets.items():
            cols = {str(c).strip().lower(): c for c in df.columns}
            for ck, uk, tk in RAW_SHEET_TYPES:
                if ck in cols and uk in cols:
                    for _, r in df.iterrows():
                        course = norm(r[cols[ck]])
                        if not course:
                            continue
                        recs.append(dict(
                            id=norm_id(r.get(cols.get("external id"))),
                            email=norm(r.get(cols.get("email"))),
                            name=r.get(cols.get("name")),
                            course=course,
                            url=None if _blank(r[cols[uk]]) else str(r[cols[uk]]).strip(),
                            time=pd.to_datetime(r.get(cols.get(tk)), errors="coerce"),
                        ))
    if not recs:
        raise ValueError("No Coursera data found: the export needs a sheet with "
                         "'Course Certificate URL' or 'Specialization Certificate URL'.")
    return recs


class RawIndex:
    def __init__(self, recs):
        self.by_id = defaultdict(list)
        self.by_email = defaultdict(list)
        self.by_name = defaultdict(list)
        self.ids, self.emails, self.names = set(), set(), set()
        self.student_courses = defaultdict(set)
        for r in recs:
            nk = name_key(r["name"])
            if r["id"]:
                self.by_id[(r["course"], r["id"])].append(r)
                self.ids.add(r["id"])
                self.student_courses[("id", r["id"])].add(r["course"])
            if r["email"]:
                self.by_email[(r["course"], r["email"])].append(r)
                self.emails.add(r["email"])
                self.student_courses[("email", r["email"])].add(r["course"])
            if nk:
                self.by_name[(r["course"], nk)].append(r)
                self.names.add(nk)


# ----------------------------------------------------------------------------- matching
def match_row(idx, rid, email, name, course):
    """Return dict describing the match for one abstract row."""
    c, nk = norm(course), name_key(name)
    by_id = idx.by_id.get((c, rid), []) if rid else []
    by_em = idx.by_email.get((c, email), []) if email else []
    cands, method = [], ""
    if by_id or by_em:
        cands = by_id + [r for r in by_em if r not in by_id]
        method = "Roll No + Email" if (by_id and by_em) else ("Roll No" if by_id else "Email")
    elif nk and idx.by_name.get((c, nk)):
        cands, method = idx.by_name[(c, nk)], "Name only"

    out = dict(method=method or "-", url=None, remark="", name_check="-", roll_check="-")
    if cands:
        with_url = [r for r in cands if r["url"]]
        ref = (with_url or cands)[0]
        # cross-checks
        if method != "Name only":
            a, b = name_key(name), name_key(ref["name"])
            if a == b:
                out["name_check"] = "OK"
            elif name_tokens(name) == name_tokens(ref["name"]):
                out["name_check"] = "OK (word order differs)"
            else:
                out["name_check"] = f"MISMATCH (export: {ref['name']})"
        else:
            out["name_check"] = "OK"
        if rid and ref["id"]:
            out["roll_check"] = "OK" if rid == ref["id"] else f"MISMATCH (export: {ref['id']})"
        elif rid and not ref["id"]:
            out["roll_check"] = "Roll No blank in export"
        if with_url:
            urls = {}
            for r in with_url:
                urls.setdefault(r["url"], r["time"])
            if len(urls) > 1:  # several different certificates -> latest completion
                best = max(urls, key=lambda u: urls[u] if not pd.isna(urls[u]) else pd.Timestamp.min)
                out["url"] = best
                out["remark"] = f"{len(urls)} different certificates found; latest used"
            else:
                out["url"] = next(iter(urls))
        else:
            out["remark"] = "Course found in export but no certificate URL issued"
        return out

    # not found -> explain why
    sc = set()
    if rid:
        sc |= idx.student_courses.get(("id", rid), set())
    if email:
        sc |= idx.student_courses.get(("email", email), set())
    known = bool(sc) or (nk in idx.names)
    if known:
        close = difflib.get_close_matches(c, list(sc), n=1, cutoff=0.85) if sc else []
        out["remark"] = ("Student found but this course is not in the export"
                         + (f" (similar: '{close[0]}')" if close else ""))
    else:
        out["remark"] = "Student not found in export"
    return out


# ----------------------------------------------------------------------------- abstract
def _find_header(ws):
    for r in range(1, min(ws.max_row, 10) + 1):
        heads = {norm(c.value): c.column for c in ws[r] if not _blank(c.value)}
        if (heads.keys() & NAME_ALIASES) and (heads.keys() & COURSE_ALIASES) and (heads.keys() & ID_ALIASES):
            return r, heads
    return None, None


def _copy_style(src, dst):
    dst.font, dst.border, dst.fill = copy(src.font), copy(src.border), copy(src.fill)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format
    dst.protection = copy(src.protection)


def process(abstract_file, raw_files, output_path):
    """Build the output workbook. Returns (stats_dataframe, issues_dataframe)."""
    idx = RawIndex(load_raw(raw_files))
    if hasattr(abstract_file, "seek"):
        abstract_file.seek(0)
    src_wb = load_workbook(abstract_file)
    out_wb = Workbook()
    out_wb.remove(out_wb.active)

    verification, course_lists, done_sheets, new_last = [], {}, [], {}
    link_font = Font(name="Calibri", size=11, color="0000FF", underline="single")
    na_fill = PatternFill("solid", fgColor="FFF2CC")

    for ws in src_wb.worksheets:
        hr, heads = _find_header(ws)
        if hr is None:
            continue  # SUMMARY / other sheets
        col = lambda aliases: next((heads[k] for k in heads if k in aliases), None)
        c_id, c_name, c_course, c_email = col(ID_ALIASES), col(NAME_ALIASES), col(COURSE_ALIASES), col(EMAIL_ALIASES)
        keep = [c for k, c in sorted(heads.items(), key=lambda x: x[1]) if k not in LINK_ALIASES]
        link_col_out = len(keep) + 1

        new = out_wb.create_sheet(ws.title)
        header_style_cell = ws.cell(hr, keep[-1])
        # header + column widths + data
        for j, sc in enumerate(keep, start=1):
            new.column_dimensions[get_column_letter(j)].width = ws.column_dimensions[get_column_letter(sc)].width or 14
            for r in range(hr, ws.max_row + 1):
                s, d = ws.cell(r, sc), new.cell(r - hr + 1, j)
                d.value = s.value
                _copy_style(s, d)
        h = new.cell(1, link_col_out, "Certificate link")
        _copy_style(header_style_cell, h)
        new.column_dimensions[get_column_letter(link_col_out)].width = 58
        new.freeze_panes = "A2"
        new.auto_filter.ref = f"A1:{get_column_letter(link_col_out)}{ws.max_row - hr + 1}"
        new.row_dimensions[1].height = ws.row_dimensions[hr].height

        courses = set()
        for r in range(hr + 1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in (c_id, c_name, c_course)]
            if all(_blank(v) for v in vals):
                continue
            rid = norm_id(ws.cell(r, c_id).value)
            email = norm(ws.cell(r, c_email).value) if c_email else ""
            m = match_row(idx, rid, email, ws.cell(r, c_name).value, ws.cell(r, c_course).value)
            out_r = r - hr + 1
            cell = new.cell(out_r, link_col_out)
            _copy_style(ws.cell(r, keep[-1]), cell)
            if m["url"]:
                cell.value, cell.hyperlink, cell.font = m["url"], m["url"], link_font
            else:
                cell.value, cell.fill = NA_TEXT, na_fill
            courses.add(str(ws.cell(r, c_course).value).strip())
            verification.append([ws.title, out_r, ws.cell(r, c_id).value, ws.cell(r, c_name).value,
                                 ws.cell(r, c_course).value, m["method"], m["name_check"],
                                 m["roll_check"], "Found" if m["url"] else NA_TEXT, m["remark"]])
        course_lists[ws.title] = (sorted(courses), keep.index(c_course) + 1, link_col_out)
        new_last[ws.title] = ws.max_row - hr + 1
        done_sheets.append(ws.title)

    if not done_sheets:
        raise ValueError("No sheet in the abstract workbook has Roll No / Name / Course headers.")

    # ---- Verification sheet
    vs = out_wb.create_sheet("Verification")
    head = ["Sheet", "Row", "Roll No", "Name", "Course", "Matched by",
            "Name check", "Roll No check", "Link status", "Remark"]
    vs.append(head)
    for c in vs[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="D9E1F2")
    for row in verification:
        vs.append(row)
    flag = PatternFill("solid", fgColor="FCE4D6")
    for r in range(2, vs.max_row + 1):
        if vs.cell(r, 9).value != "Found" or "MISMATCH" in str(vs.cell(r, 7).value) \
                or "MISMATCH" in str(vs.cell(r, 8).value):
            for c in vs[r]:
                c.fill = flag
    for i, w in enumerate([24, 6, 12, 30, 52, 16, 28, 24, 18, 60], 1):
        vs.column_dimensions[get_column_letter(i)].width = w
    vs.freeze_panes = "A2"
    vs.auto_filter.ref = f"A1:J{vs.max_row}"

    # ---- SUMMARY sheet (live formulas)
    sm = out_wb.create_sheet("SUMMARY")
    row = 1
    for title in done_sheets:
        courses, c_course, lc = course_lists[title]
        cl, ll = get_column_letter(c_course), get_column_letter(lc)
        q = f"'{title}'"
        last = new_last[title]
        CR, LR = f"{q}!${cl}$2:${cl}${last}", f"{q}!${ll}$2:${ll}${last}"
        sm.cell(row, 1, f"{title} - course-wise summary").font = Font(bold=True, size=12)
        row += 1
        for j, t in enumerate(["Course", "Students", "Links found", NA_TEXT], 1):
            c = sm.cell(row, j, t)
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="D9E1F2")
            c.alignment = Alignment(horizontal="center")
        first = row + 1
        for course in courses:
            row += 1
            sm.cell(row, 1, course)
            sm.cell(row, 2, f'=SUMPRODUCT(--(TRIM({CR})=A{row}))')
            sm.cell(row, 3, f'=SUMPRODUCT(--(TRIM({CR})=A{row}),--(LEFT({LR},4)="http"))')
            sm.cell(row, 4, f'=SUMPRODUCT(--(TRIM({CR})=A{row}),--({LR}="{NA_TEXT}"))')
        row += 1
        sm.cell(row, 1, "TOTAL").font = Font(bold=True)
        for j in (2, 3, 4):
            L = get_column_letter(j)
            c = sm.cell(row, j, f"=SUM({L}{first}:{L}{row - 1})")
            c.font = Font(bold=True)
        row += 3
    for i, w in enumerate([58, 12, 14, 18], 1):
        sm.column_dimensions[get_column_letter(i)].width = w

    out_wb.save(output_path)

    df = pd.DataFrame(verification, columns=head)
    issues = df[(df["Link status"] != "Found") | df["Name check"].str.contains("MISMATCH")
                | df["Roll No check"].str.contains("MISMATCH")]
    return df, issues


# ----------------------------------------------------------------------------- CLI
if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Paste Coursera certificate links into the abstract workbook.")
    ap.add_argument("abstract")
    ap.add_argument("exports", nargs="+")
    ap.add_argument("-o", "--output", default="Abstract_with_certificate_links.xlsx")
    a = ap.parse_args()
    allrows, issues = process(a.abstract, a.exports, a.output)
    print(f"Done -> {a.output}\nRows: {len(allrows)} | Links found: {(allrows['Link status'] == 'Found').sum()} "
          f"| Not available: {(allrows['Link status'] != 'Found').sum()}")
